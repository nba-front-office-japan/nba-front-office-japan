# 「NBA_Payroll_2026-27_to_2031-32.xlsx」から、サラリーページ用のデータ(lib/salary/payroll-data.ts)を生成する。
# Excelファイルそのものはリポジトリにもサイトにも置かない(数値だけをTypeScriptのデータにして取り込む)。
#
# 使い方: python scripts/build-salary-data.py "<Excelのパス>"
#   既定のパス: C:/Users/user/OneDrive/デスクトップ/NBA DATA BASE/NBA_Payroll_2026-27_to_2031-32.xlsx
#
# 読み取る内容:
#   All Teams Summary : 30チームの年度別総年俸(5行目が見出し: Rk, Team, 2026-27 … 2031-32)
#   各チームのシート   : 5行目が見出し(Player, Age, 2026-27 … 2031-32, Guaranteed, Options, Option Source)、
#                        6行目から選手、最後が Team Totals
import io
import json
import re
import sys

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "C:/Users/user/OneDrive/デスクトップ/NBA DATA BASE/NBA_Payroll_2026-27_to_2031-32.xlsx"
OUT = "lib/salary/payroll-data.ts"
SEASONS = ["2026-27", "2027-28", "2028-29", "2029-30", "2030-31", "2031-32"]

# Excelのチーム名 → 当サイトのチーム略称(lib/team-colors.ts・teams.abbreviation と同じ)
ABBR = {
    "Atlanta Hawks": "ATL", "Boston Celtics": "BOS", "Brooklyn Nets": "BKN", "Charlotte Hornets": "CHA",
    "Chicago Bulls": "CHI", "Cleveland Cavaliers": "CLE", "Dallas Mavericks": "DAL", "Denver Nuggets": "DEN",
    "Detroit Pistons": "DET", "Golden State Warriors": "GSW", "Houston Rockets": "HOU", "Indiana Pacers": "IND",
    "Los Angeles Clippers": "LAC", "Los Angeles Lakers": "LAL", "Memphis Grizzlies": "MEM", "Miami Heat": "MIA",
    "Milwaukee Bucks": "MIL", "Minnesota Timberwolves": "MIN", "New Orleans Pelicans": "NOP", "New York Knicks": "NYK",
    "Oklahoma City Thunder": "OKC", "Orlando Magic": "ORL", "Philadelphia 76ers": "PHI", "Phoenix Suns": "PHX",
    "Portland Trail Blazers": "POR", "Sacramento Kings": "SAC", "San Antonio Spurs": "SAS", "Toronto Raptors": "TOR",
    "Utah Jazz": "UTA", "Washington Wizards": "WAS",
}

HEADER = ("Player", "Age", *SEASONS, "Guaranteed", "Options", "Option Source")


def money(v, where):
    if v is None:
        return None
    if not isinstance(v, (int, float)) or v < 0 or v != int(v):
        raise SystemExit(f"金額が正しくありません: {where}: {v!r}")
    return int(v)


def parse_options(text, where):
    if not text:
        return []
    out = []
    for part in str(text).split(";"):
        m = re.fullmatch(r"\s*(\d{4}-\d{2}): (Player|Team) Option\s*", part)
        if not m or m.group(1) not in SEASONS:
            raise SystemExit(f"オプションの書式が想定外です: {where}: {text!r}")
        out.append({"season": m.group(1), "type": "player" if m.group(2) == "Player" else "team"})
    return out


wb = openpyxl.load_workbook(SRC, data_only=True)

# --- All Teams Summary ---
ws = wb["All Teams Summary"]
rows = list(ws.iter_rows(values_only=True))
assert rows[4][:8] == ("Rk", "Team", *SEASONS), rows[4]
teams = []
for r in rows[5:]:
    if r[0] is None:
        continue
    name = r[1]
    if name not in ABBR:
        raise SystemExit(f"チーム名が想定外です: {name}")
    teams.append({"abbr": ABBR[name], "name": name, "totals": [money(v, f"{name} 総年俸") for v in r[2:8]]})
assert len(teams) == 30, len(teams)
assert len({t["abbr"] for t in teams}) == 30

# --- 各チームのシート ---
players = []
checks = []
for t in teams:
    ws = wb[t["name"]]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[4][:11] == HEADER, (t["name"], rows[4])
    team_totals = None
    team_players = []
    for i, r in enumerate(rows[5:], start=6):
        if r[0] is None:
            if any(v is not None for v in r[:11]):
                raise SystemExit(f"選手名のない行があります: {t['name']} {i}行目")
            continue
        if r[0] == "Team Totals":
            team_totals = [money(v, f"{t['name']} Team Totals") for v in r[2:8]]
            continue
        where = f"{t['name']} {i}行目 {r[0]}"
        age = r[1]
        if age is not None and (not isinstance(age, (int, float)) or not 18 <= age <= 45):
            raise SystemExit(f"年齢が想定外です: {where}: {age!r}")
        team_players.append({
            "name": str(r[0]).strip(),
            "team": t["abbr"],
            "age": int(age) if age is not None else None,
            "salaries": [money(v, where) for v in r[2:8]],
            "guaranteed": money(r[8], where),
            "options": parse_options(r[9], where),
        })
    if team_totals is None:
        raise SystemExit(f"Team Totals の行がありません: {t['name']}")
    sums = [sum(p["salaries"][k] or 0 for p in team_players) for k in range(6)]
    sheet = [v or 0 for v in team_totals]
    if sums != sheet:
        raise SystemExit(f"選手の合計とシートの Team Totals が一致しません: {t['name']}")
    summary = [v or 0 for v in t["totals"]]
    if sums != summary:
        checks.append((t["name"], [s - v for s, v in zip(sums, summary)]))
    # Excelの順(2026-27の年俸が高い順)を保つ
    players.extend(team_players)

print("teams:", len(teams), "players:", len(players))
print("選手の合計と All Teams Summary の差:", checks or "なし")

ts = (
    "// 「NBA Payroll 2026-27–2031-32」(運営者が用意したExcel)から scripts/build-salary-data.py で生成したデータ。\n"
    "// 手で編集せず、Excelを更新したら生成し直す。Excelファイルそのものはサイトに置かない。\n\n"
    "import type { SalaryPlayer, SalaryTeam } from \"./types\";\n\n"
    f"export const SALARY_TEAMS: SalaryTeam[] = {json.dumps(teams, ensure_ascii=False)};\n\n"
    f"export const SALARY_PLAYERS: SalaryPlayer[] = {json.dumps(players, ensure_ascii=False)};\n"
)
# 1行が長くなりすぎないよう、要素ごとに改行する
ts = ts.replace("}, {", "},\n  {").replace("= [{", "= [\n  {").replace("}];", "},\n];")
io.open(OUT, "w", encoding="utf-8", newline="\n").write(ts)
print("wrote", OUT)
